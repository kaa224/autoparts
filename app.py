#!/usr/bin/env python3
"""Web interface for batch generation of Ozon product cards."""

from __future__ import annotations

import base64
import hmac
import json
import os
import re
import threading
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import BinaryIO

import xlrd
from flask import Flask, abort, jsonify, render_template, request, send_file, send_from_directory, url_for
from openpyxl import load_workbook
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.utils import secure_filename

from generate_card import build_with_name, normalize_article


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("AUTOPARTS_DATA_DIR", BASE_DIR / "data"))
TEMPLATE_PATH = Path(os.environ.get("AUTOPARTS_TEMPLATE", BASE_DIR / "assets" / "Frame 6.png"))
ALLOWED_EXTENSIONS = {".xls", ".xlsx"}
ARTICLE_FILENAME = re.compile(r"[^0-9A-Za-zА-Яа-я._-]+")
JOB_ID = re.compile(r"^[0-9a-f]{12}$")

app = Flask(__name__)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1, x_prefix=1)
app.config["MAX_CONTENT_LENGTH"] = 15 * 1024 * 1024
executor = ThreadPoolExecutor(max_workers=int(os.environ.get("AUTOPARTS_WORKERS", "2")))
jobs_lock = threading.Lock()


@dataclass
class Row:
    article: str
    name: str
    status: str = "waiting"
    error: str | None = None
    filename: str | None = None
    models_filename: str | None = None
    json_filename: str | None = None

    def json(self, job_id: str) -> dict[str, str | None]:
        return {
            "article": self.article,
            "name": self.name,
            "status": self.status,
            "error": self.error,
            "file_url": url_for("job_file", job_id=job_id, filename=self.filename)
            if self.filename else None,
            "json_url": url_for("job_file", job_id=job_id, filename=self.json_filename)
            if self.json_filename else None,
            "models_url": url_for("job_file", job_id=job_id, filename=self.models_filename)
            if self.models_filename else None,
        }


@dataclass
class Job:
    id: str
    rows: list[Row]
    output_dir: Path
    created_at: str = field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    status: str = "waiting"
    source_filename: str = ""

    def json(self) -> dict:
        complete = sum(row.status in {"done", "error"} for row in self.rows)
        return {
            "id": self.id,
            "status": self.status,
            "complete": complete,
            "total": len(self.rows),
            "rows": [row.json(self.id) for row in self.rows],
            "folder_url": url_for("job_files", job_id=self.id),
            "download_url": url_for("download_job", job_id=self.id),
            "status_url": url_for("job_status", job_id=self.id),
            "local_path": str(self.output_dir.resolve()),
        }


jobs: dict[str, Job] = {}


def persist_job(job: Job) -> None:
    job_dir = job.output_dir.parent
    job_dir.mkdir(parents=True, exist_ok=True)
    manifest = {
        "id": job.id,
        "created_at": job.created_at,
        "source_filename": job.source_filename,
        "status": job.status,
        "total": len(job.rows),
        "done": sum(row.status == "done" for row in job.rows),
        "errors": sum(row.status == "error" for row in job.rows),
    }
    (job_dir / "job.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def stored_job(job_id: str) -> Job | None:
    if not JOB_ID.fullmatch(job_id):
        return None
    job_dir = DATA_DIR / "jobs" / job_id
    output_dir = job_dir / "output"
    if not output_dir.is_dir():
        return None
    manifest_path = job_dir / "job.json"
    manifest: dict = {}
    if manifest_path.is_file():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            manifest = {}
    fallback = datetime.fromtimestamp(job_dir.stat().st_ctime, timezone.utc).isoformat()
    return Job(
        id=job_id,
        rows=[],
        output_dir=output_dir,
        created_at=manifest.get("created_at", fallback),
        status=manifest.get("status", "done"),
        source_filename=manifest.get("source_filename", ""),
    )


def archive_groups() -> list[dict]:
    root = DATA_DIR / "jobs"
    groups: dict[str, list[dict]] = {}
    if not root.is_dir():
        return []
    for job_dir in root.iterdir():
        job = stored_job(job_dir.name)
        if job is None:
            continue
        try:
            created = datetime.fromisoformat(job.created_at)
        except ValueError:
            created = datetime.fromtimestamp(job_dir.stat().st_ctime, timezone.utc)
        artifacts = sorted(
            path for path in job.output_dir.iterdir()
            if path.suffix.lower() in {".jpg", ".json"}
        )
        item = {
            "job": job,
            "created": created,
            "time": created.astimezone().strftime("%H:%M"),
            "file_count": len(artifacts),
            "card_count": sum(path.suffix.lower() == ".jpg" for path in artifacts),
            "size": sum(path.stat().st_size for path in artifacts),
        }
        groups.setdefault(created.date().isoformat(), []).append(item)
    result = []
    for date_key in sorted(groups, reverse=True):
        items = sorted(groups[date_key], key=lambda item: item["created"], reverse=True)
        date_value = datetime.fromisoformat(date_key)
        result.append({"date": date_key, "label": date_value.strftime("%d.%m.%Y"), "items": items})
    return result


def unauthorized_response():
    return (
        "Требуется авторизация\n",
        401,
        {"WWW-Authenticate": 'Basic realm="Ozon card generator", charset="UTF-8"'},
    )


@app.before_request
def require_separate_auth():
    """Protect this application with credentials independent of rialsat-admin."""
    if os.environ.get("AUTOPARTS_REQUIRE_AUTH", "0") != "1":
        return None
    expected_login = os.environ.get("AUTOPARTS_LOGIN", "")
    expected_password = os.environ.get("AUTOPARTS_PASSWORD", "")
    if not expected_login or not expected_password:
        return "Авторизация приложения не настроена\n", 503
    header = request.headers.get("Authorization", "")
    scheme, _, token = header.partition(" ")
    if scheme.casefold() != "basic" or not token:
        return unauthorized_response()
    try:
        decoded = base64.b64decode(token, validate=True).decode("utf-8")
    except (ValueError, UnicodeDecodeError):
        return unauthorized_response()
    login, separator, password = decoded.partition(":")
    if not separator:
        return unauthorized_response()
    login_ok = hmac.compare_digest(login, expected_login)
    password_ok = hmac.compare_digest(password, expected_password)
    if not (login_ok and password_ok):
        return unauthorized_response()
    return None


def normalized_header(value: object) -> str:
    return " ".join(str(value or "").strip().casefold().split())


def rows_from_xls(stream: BinaryIO) -> list[tuple[str, str]]:
    book = xlrd.open_workbook(file_contents=stream.read())
    sheet = book.sheet_by_index(0)
    values = [[sheet.cell_value(row, col) for col in range(sheet.ncols)] for row in range(sheet.nrows)]
    return parse_table(values)


def rows_from_xlsx(stream: BinaryIO) -> list[tuple[str, str]]:
    book = load_workbook(stream, read_only=True, data_only=True)
    sheet = book.active
    return parse_table([list(row) for row in sheet.iter_rows(values_only=True)])


def parse_table(values: list[list[object]]) -> list[tuple[str, str]]:
    if not values:
        raise ValueError("Файл пуст")
    headers = [normalized_header(value) for value in values[0]]
    missing = [name for name in ("артикул", "наименование") if name not in headers]
    if missing:
        raise ValueError("Нет обязательных столбцов: " + ", ".join(missing))
    article_col, name_col = headers.index("артикул"), headers.index("наименование")
    result: list[tuple[str, str]] = []
    for source_row in values[1:]:
        article = normalize_article(source_row[article_col] if article_col < len(source_row) else "")
        name = " ".join(str(source_row[name_col] if name_col < len(source_row) else "").split())
        if article and name:
            result.append((article, name))
    if not result:
        raise ValueError("В файле нет строк с заполненными артикулом и наименованием")
    return result


def safe_article_filename(article: str) -> str:
    cleaned = ARTICLE_FILENAME.sub("_", article).strip("._")
    return cleaned or "item"


def process_job(job_id: str) -> None:
    job = jobs[job_id]
    with jobs_lock:
        job.status = "processing"
        persist_job(job)
    for row in job.rows:
        with jobs_lock:
            row.status = "processing"
        try:
            file_article = safe_article_filename(row.article)
            result = build_with_name(
                row.article, row.name, TEMPLATE_PATH, job.output_dir,
                output_stem=file_article,
            )
            expected = job.output_dir / f"{file_article}.jpg"
            if result != expected:
                result.replace(expected)
            with jobs_lock:
                row.filename = expected.name
                row.models_filename = f"{file_article}_models.jpg"
                row.json_filename = f"{file_article}.json"
                row.status = "done"
        except Exception as exc:  # one failed article must not stop the batch
            with jobs_lock:
                row.status = "error"
                row.error = str(exc)
    with jobs_lock:
        job.status = "done"
        persist_job(job)


def get_job(job_id: str) -> Job:
    job = jobs.get(job_id) or stored_job(job_id)
    if job is None:
        abort(404)
    return job


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/archives")
def archives():
    return render_template("archives.html", groups=archive_groups())


@app.post("/api/jobs")
def create_job():
    upload = request.files.get("file")
    if upload is None or not upload.filename:
        return jsonify(error="Выберите Excel-файл"), 400
    suffix = Path(secure_filename(upload.filename)).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        return jsonify(error="Поддерживаются только файлы .xls и .xlsx"), 400
    try:
        parsed = rows_from_xls(upload.stream) if suffix == ".xls" else rows_from_xlsx(upload.stream)
    except (ValueError, xlrd.XLRDError, OSError) as exc:
        return jsonify(error=f"Не удалось прочитать файл: {exc}"), 400

    job_id = uuid.uuid4().hex[:12]
    output_dir = DATA_DIR / "jobs" / job_id / "output"
    job = Job(
        job_id,
        [Row(article, name) for article, name in parsed],
        output_dir,
        source_filename=secure_filename(upload.filename),
    )
    with jobs_lock:
        jobs[job_id] = job
        persist_job(job)
    executor.submit(process_job, job_id)
    return jsonify(job.json()), 202


@app.get("/api/jobs/<job_id>")
def job_status(job_id: str):
    with jobs_lock:
        payload = get_job(job_id).json()
    return jsonify(payload)


@app.get("/jobs/<job_id>/files")
def job_files(job_id: str):
    job = get_job(job_id)
    files = sorted(
        path for path in job.output_dir.iterdir() if path.suffix.lower() in {".jpg", ".json"}
    ) if job.output_dir.exists() else []
    return render_template("files.html", job=job, files=files)


@app.get("/jobs/<job_id>/files/<path:filename>")
def job_file(job_id: str, filename: str):
    job = get_job(job_id)
    return send_from_directory(job.output_dir, filename, as_attachment=True)


@app.get("/jobs/<job_id>/download")
def download_job(job_id: str):
    job = get_job(job_id)
    job.output_dir.mkdir(parents=True, exist_ok=True)
    try:
        date_prefix = datetime.fromisoformat(job.created_at).date().isoformat()
    except ValueError:
        date_prefix = "archive"
    archive = job.output_dir.parent / f"cards-{date_prefix}-{job_id}.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
        for artifact in sorted(job.output_dir.iterdir()):
            if artifact.suffix.lower() in {".jpg", ".json"}:
                bundle.write(artifact, artifact.name)
    return send_file(archive, as_attachment=True, download_name=archive.name)


@app.get("/health")
def health():
    return {"status": "ok", "template": TEMPLATE_PATH.exists()}


if __name__ == "__main__":
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "8000")), debug=False)
